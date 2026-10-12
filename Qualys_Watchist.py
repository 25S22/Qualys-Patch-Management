"""
qualys_cve_watchlist.py
=======================
Pulls EVERY critical vulnerability (Qualys severity 4 and 5) your estate is
affected by, and writes an Excel file with two views that can be used as
context for building a Vulnerability Watchlist in Recorded Future:

   1. All CVEs          – every CVE, ranked by the number of assets affected
   2. By Application    – CVEs clubbed per application, with the top 10 CVEs
                          of each application

APIs USED  (VMDR only — nothing else is called)
───────────────────────────────────────────────
  1. POST /api/2.0/fo/session/                  login / logout
  2. GET  /api/2.0/fo/asset/host/vm/detection/  Host List VM Detection
  3. POST /api/2.0/fo/knowledge_base/vuln/      VM Knowledge Base

  No other Qualys module or API is called. Asset counts come purely from
  VMDR host detections.

HOW IT WORKS
────────────
 Step 1 — FETCH DETECTIONS
   Reads host detections filtered by severity and status (default: severity
   4-5, New/Active/Re-Opened) and follows the pagination links until every
   host has been read.  Builds a map of  QID → {set of host IDs}.
   Potential detections are dropped client-side unless INCLUDE_POTENTIAL=True.

 Step 2 — ENRICH
   Looks up every QID in the Knowledge Base for its CVE list, title, CVSS,
   threat-intelligence tags, vendor/product and publish date.

 Step 3 — BUILD THE TWO VIEWS
   Asset counts are always UNIQUE hosts (a host is never counted twice).

   All CVEs:  one row per CVE; assets = unique hosts across every QID that
              references the CVE.

   By Application:  QIDs are clubbed into one application by normalising the
              Qualys title, ignoring version / bulletin / month, e.g.
                 Google Chrome Prior to 126.0.x Multiple Vulnerabilities
                 Google Chrome Prior to 127.0.x Multiple Vulnerabilities
                                                  → "Google Chrome"
              Inside each application the CVEs are ranked
                 CISA KEV → actively attacked → CVSS v3 → assets affected
              and the first CVES_PER_APP (default 10) are kept.

 Step 4 — EXPORT
   Saves one Excel workbook:
     "All CVEs"          – every CVE found
     "By Application"    – one row per application, top CVEs listed in a cell
     "Top CVEs per App"  – the same top CVEs as flat rows (easy to filter)
     "QID Mapping"       – which QID/title landed in which application (use it
                           to check the grouping and tune PRODUCT_OVERRIDES)
     "Run Info"          – parameters and totals for this run

   This script is READ-ONLY — it never changes anything in Qualys.

TROUBLESHOOTING
───────────────
  • Errors print Qualys's own message (HTTP status, code and text).
  • If your platform rejects an optional parameter ("Unrecognized
    parameter(s): …") the script drops it automatically and retries.
  • If nothing is found, the script runs an unfiltered test call and tells you
    whether the problem is your filters or your account's asset scope.
  • Two applications merged, or one split in two?  Open "QID Mapping", then
    add a rule to PRODUCT_OVERRIDES.
  • Set DEBUG_SAVE_XML = True to keep the raw first-page XML for inspection.

PERMISSIONS REQUIRED
────────────────────
  • API Access
  • VMDR access to the assets you want counted (counts only reflect hosts your
    account's asset groups / tags can see)

DEPENDENCIES
────────────
  pip install requests pandas openpyxl
"""

import os
import re
import sys
import time
import logging
import xml.etree.ElementTree as ET
from collections import defaultdict
from datetime import datetime, timezone

import requests
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ──────────────────────────────────────────────────────────────────────────────
# CONFIGURATION  ←  edit these before running
# ──────────────────────────────────────────────────────────────────────────────
# Credentials: environment variables win if set, otherwise the defaults below.
USERNAME     = os.getenv("QUALYS_USERNAME", "your_qualys_username")
PASSWORD     = os.getenv("QUALYS_PASSWORD", "your_qualys_password")

# Your Qualys API server URL  (e.g. https://qualysapi.qg1.apps.qualys.in)
# Find it: Qualys UI → Help → About → Security Operations Center
BASE_URL     = "https://qualysapi.qg1.apps.qualys.in"

# SSL certificate bundle.  Use a .pem path, True (system CAs), or False for a
# trusted test environment only.
CERT_PATH    = "/path/to/your/corporate_cert.pem"

# ── What to report ───────────────────────────────────────────────────────────
# Qualys severity levels to include.  5 = Urgent, 4 = Critical, 3 = Serious
# Accepts a single level, a comma list, or a range:  "5"  |  "4,5"  |  "4-5"
SEVERITY_LEVELS  = "4-5"

# How many CVEs to list per application in the "By Application" views.
# (ranked KEV → actively attacked → CVSS v3 → assets).  0 = list all.
CVES_PER_APP     = 10

# Optional noise filter: only keep CVEs affecting MORE THAN this many hosts.
# 0 = keep everything (pull ALL CVEs).
MIN_ASSET_COUNT  = 0

# Optional manual grouping rules, checked BEFORE the automatic ones.
# (regex matched against the Qualys title, case-insensitive; first match wins)
# Examples:
#   (r"^Google Chrome",              "Google Chrome"),
#   (r"Microsoft Edge",              "Microsoft Edge"),
#   (r"Mozilla Firefox( ESR)?",      "Mozilla Firefox"),
#   (r"Windows.*Security Update",    "Microsoft Windows"),
PRODUCT_OVERRIDES = [
]

# Detection status to count.  Fixed detections are excluded by default.
DETECTION_STATUS = "New,Active,Re-Opened"

# False = count only CONFIRMED detections (recommended, fewer false positives)
# True  = also count POTENTIAL detections
INCLUDE_POTENTIAL = False

# Optional extra filters passed straight to the detection API, e.g.
#   {"ips": "10.0.0.0/16"}
#   {"use_tags": 1, "tag_set_by": "name", "tag_set_include": "Production"}
# Strongly recommended on very large subscriptions to keep the pull quick.
EXTRA_DETECTION_PARAMS = {}

# ── Performance / reliability ────────────────────────────────────────────────
HOSTS_PER_PAGE   = 1000    # hosts returned per detection page (truncation_limit)
KB_BATCH_SIZE    = 100     # QIDs per Knowledge Base request
TIMEOUT          = 600     # seconds to wait for a single API response
MAX_RETRIES      = 5       # attempts per request on 409 / 429 / 5xx
RETRY_WAIT       = 15      # seconds to wait if Qualys doesn't send Retry-After

# True = save raw first-page XML files (qualys_debug_*.xml) for inspection
DEBUG_SAVE_XML   = False

# ── Output ───────────────────────────────────────────────────────────────────
# A timestamp is inserted before the extension, e.g. qualys_cve_watchlist_20250101_120000.xlsx
OUTPUT_FILE      = "qualys_cve_watchlist.xlsx"
# ──────────────────────────────────────────────────────────────────────────────


# ── Logging (stdout only — clean terminal output) ─────────────────────────────
logging.basicConfig(
    stream=sys.stdout,
    level=logging.INFO,
    format="%(message)s",
)
log = logging.getLogger("qualys_cve_watchlist")

LINE  = "─" * 65
DLINE = "═" * 65
EXCEL_CELL_LIMIT = 32000      # Excel allows 32,767 characters per cell


# ── Small helpers ─────────────────────────────────────────────────────────────
class QualysAPIError(RuntimeError):
    """Raised for any Qualys error, carrying the real message."""

    def __init__(self, status, code, text):
        self.status, self.code, self.text = status, code, text
        super().__init__(f"Qualys API error — HTTP {status}"
                         f"{' | code ' + code if code else ''}: {text}")


def _to_float(val):
    try:
        return float(val)
    except (TypeError, ValueError):
        return None


def _max_none(a, b):
    """max() that ignores None."""
    if a is None:
        return b
    if b is None:
        return a
    return max(a, b)


def _clip(text: str) -> str:
    return text if len(text) <= EXCEL_CELL_LIMIT else text[:EXCEL_CELL_LIMIT] + " …"


def _debug_save(name: str, content: bytes):
    if DEBUG_SAVE_XML:
        path = f"qualys_debug_{name}.xml"
        with open(path, "wb") as fh:
            fh.write(content)
        log.info(f"  [debug] raw response saved → {path}")


# ── Title → application normalisation ─────────────────────────────────────────
# Linux distro advisories are titled "<Distro> Update for <package> (ID)";
# for these the PACKAGE is the application, e.g. "Red Hat kernel".
_DISTRO_START = re.compile(
    r"^(Red\s*Hat(?:\s+Enterprise\s+Linux)?|Ubuntu|Debian|CentOS|"
    r"SUSE(?:\s+Enterprise\s+Linux)?|openSUSE|Fedora|Amazon\s+Linux|AlmaLinux|"
    r"Rocky\s+Linux|Oracle\s+(?:Enterprise\s+)?Linux|Alibaba\s+Cloud\s+Linux|"
    r"EulerOS|Photon\s+OS)\b", re.I)

# Words/phrases that mark the END of the product name in a Qualys title
_CUT = re.compile(
    r"\s+(?:prior\s+to|before|earlier\s+than|up\s+to|"
    r"security\s+(?:update|advisory|bulletin|fix|patch|release|notification)s?|"
    r"critical\s+patch\s+update|cumulative|update|patch|"
    r"multiple|vulnerabilit(?:y|ies)|remote\s+code|denial\s+of|privilege|"
    r"information\s+disclosure|elevation\s+of|security\s+feature|arbitrary|"
    r"authentication\s+bypass|command\s+injection|sql\s+injection|cross[- ]site|"
    r"path\s+traversal|heap|stack|buffer|integer|out[- ]of[- ]bounds|"
    r"use[- ]after[- ]free|memory\s+corruption|zero[- ]day|improper|"
    r"insufficient|unauthorized|is\s+vulnerable|end\s+of\s+life)\b", re.I)

# Version numbers: "126.0.6478.182", "v9.0", "8u301", "Version 1809"
_VERSION = re.compile(r"\s+(?:version\s+)?v?\d+(?:\.\d+)+|\s+\d+u\d+\b|\s+version\b", re.I)


def normalize_title(title: str) -> str:
    """
    'Google Chrome Prior to 126.0.6478.182 Multiple Vulnerabilities' → 'Google Chrome'
    'Microsoft Windows Security Update for June 2024'                  → 'Microsoft Windows'
    'Red Hat Update for kernel (RHSA-2024:1234)'                      → 'Red Hat kernel'
    'Apache Tomcat 9.0.0.M1 to 9.0.89 Multiple Vulnerabilities'       → 'Apache Tomcat'
    """
    original = (title or "").strip()
    t = re.sub(r"\([^)]*\)", " ", original)          # drop (advisory IDs / platforms)
    t = re.sub(r"\s+", " ", t).strip()

    m = _DISTRO_START.match(t)
    if m:
        pk = re.search(r"\bfor\s+(.+)$", t[m.end():], re.I)
        if pk:
            pkg = re.sub(r"\s*:.*$", "", pk.group(1))     # "kernel : ALAS-2024-1"
            pkg = _CUT.split(pkg, maxsplit=1)[0]
            pkg = _VERSION.split(pkg, maxsplit=1)[0].strip(" -–:,.")
            if pkg:
                return f"{m.group(1)} {pkg}"

    t = _CUT.split(t, maxsplit=1)[0]
    t = _VERSION.split(t, maxsplit=1)[0].strip(" -–:,.")
    return t if len(t) >= 3 else original


def group_key(title: str) -> str:
    """Return the application a Qualys title belongs to."""
    title = (title or "").strip()
    for pattern, name in PRODUCT_OVERRIDES:
        if re.search(pattern, title, re.I):
            return name
    return normalize_title(title) or "(untitled)"


# ──────────────────────────────────────────────────────────────────────────────
class QualysCVEWatchlist:

    def __init__(self):
        self.base_url   = BASE_URL.rstrip("/")
        self.session    = requests.Session()
        self.fo_headers = {"X-Requested-With": "QualysCVEWatchlist"}
        self.det_url    = f"{self.base_url}/api/2.0/fo/asset/host/vm/detection/"

    # ── Config sanity check ──────────────────────────────────────────────────
    @staticmethod
    def _check_config():
        problems = []
        if USERNAME.startswith("your_") or PASSWORD.startswith("your_"):
            problems.append(
                "USERNAME / PASSWORD still hold placeholder values "
                "(edit them, or set QUALYS_USERNAME / QUALYS_PASSWORD)."
            )
        if isinstance(CERT_PATH, str) and not os.path.isfile(CERT_PATH):
            problems.append(
                f"CERT_PATH file not found: {CERT_PATH}  "
                "(use a real .pem path, True, or False for testing)."
            )
        if not BASE_URL.startswith("https://"):
            problems.append("BASE_URL must start with https://")
        for i, rule in enumerate(PRODUCT_OVERRIDES):
            try:
                if len(rule) != 2:
                    raise ValueError
                re.compile(rule[0])
            except (re.error, ValueError, TypeError, IndexError):
                problems.append(f"PRODUCT_OVERRIDES entry #{i+1} is invalid "
                                f"(expected (regex, name)): {rule!r}")
        if problems:
            raise ValueError("Configuration problem(s):\n    - " + "\n    - ".join(problems))

    # ── Session ──────────────────────────────────────────────────────────────
    def login(self):
        self._request(
            "POST", f"{self.base_url}/api/2.0/fo/session/",
            data={"action": "login", "username": USERNAME, "password": PASSWORD},
        )
        if "QualysSession" not in self.session.cookies:
            raise RuntimeError(
                "Login failed — no session cookie returned.\n"
                "Check USERNAME, PASSWORD and BASE_URL."
            )
        log.info(f"  ✔  Logged in to {self.base_url}")

    def logout(self):
        try:
            self.session.post(
                f"{self.base_url}/api/2.0/fo/session/",
                headers=self.fo_headers,
                data={"action": "logout"},
                verify=CERT_PATH,
                timeout=60,
            )
        except Exception:
            pass
        log.info("  ✔  Logged out")

    # ── HTTP helpers ─────────────────────────────────────────────────────────
    @staticmethod
    def _api_error(r: requests.Response) -> QualysAPIError:
        """Build an error that includes whatever Qualys said."""
        code = text = ""
        try:
            root = ET.fromstring(r.content)
            code = root.findtext(".//CODE", default="").strip()
            text = root.findtext(".//TEXT", default="").strip()
        except ET.ParseError:
            pass
        if not text:
            text = (r.text or "").strip()[:500] or "(empty response body)"
        if r.status_code == 401:
            text += "  → check credentials; the account needs API Access."
        elif r.status_code == 403:
            text += "  → permission denied; check API Access, role and asset-group scope."
        elif r.status_code in (409, 429):
            text += "  → concurrency / rate limit; wait a few minutes and re-run."
        return QualysAPIError(r.status_code, code, text)

    def _request(self, method: str, url: str, **kwargs) -> requests.Response:
        """
        Send a request with automatic retry on 409 / 429 / 5xx.
        Qualys uses 409 for concurrency / rate-limit responses and sends a
        Retry-After header — we honour it when present.
        """
        for attempt in range(1, MAX_RETRIES + 1):
            r = self.session.request(
                method,
                url,
                headers=self.fo_headers,
                verify=CERT_PATH,
                timeout=TIMEOUT,
                **kwargs,
            )
            retryable = r.status_code in (409, 429) or r.status_code >= 500
            if retryable and attempt < MAX_RETRIES:
                try:
                    wait = int(r.headers.get("Retry-After", RETRY_WAIT))
                except ValueError:
                    wait = RETRY_WAIT
                log.info(
                    f"  ⚠  HTTP {r.status_code} — waiting {wait}s and retrying "
                    f"({attempt}/{MAX_RETRIES})…"
                )
                time.sleep(wait)
                continue
            if r.status_code >= 400:
                raise self._api_error(r)
            return r
        raise RuntimeError("unreachable")   # pragma: no cover

    @staticmethod
    def _parse(r: requests.Response) -> ET.Element:
        """Parse XML and raise if Qualys returned a SIMPLE_RETURN error."""
        try:
            root = ET.fromstring(r.content)
        except ET.ParseError:
            snippet = (r.text or "")[:300].replace("\n", " ")
            raise RuntimeError(f"Qualys returned a non-XML response: {snippet}")
        if root.tag == "SIMPLE_RETURN":
            code = root.findtext(".//CODE", default="").strip()
            text = root.findtext(".//TEXT", default="").strip()
            raise QualysAPIError(r.status_code, code, text)
        return root

    @staticmethod
    def _unrecognized_params(text: str) -> list:
        """Pull parameter names out of 'Unrecognized parameter(s): a, b (…)'."""
        m = re.search(r"Unrecognized parameter\(s\):\s*([^()]+)", text or "", re.I)
        return [p.strip() for p in m.group(1).split(",") if p.strip()] if m else []

    # ── Step 1: detections → QID → hosts ─────────────────────────────────────
    def fetch_detections(self) -> tuple[dict, int, int]:
        """
        Returns:
          qid_hosts        : { "qid": {host_id, host_id, …} }
          hosts_returned   : number of hosts Qualys sent back (before client filters)
          skipped_potential: detections dropped because INCLUDE_POTENTIAL is False
        """
        params = {
            "action":           "list",
            "severities":       SEVERITY_LEVELS,
            "status":           DETECTION_STATUS,
            "show_results":     0,          # skip scan output → much smaller payload
            "truncation_limit": HOSTS_PER_PAGE,
            "output_format":    "XML",
        }
        params.update(EXTRA_DETECTION_PARAMS)

        url               = self.det_url
        qid_hosts         = defaultdict(set)
        hosts_seen        = set()
        skipped_potential = 0
        page              = 0

        while url:
            page += 1
            try:
                r    = self._request("GET", url, params=params if page == 1 else None)
                root = self._parse(r)
            except QualysAPIError as e:
                # Platform rejected an optional parameter → drop it and retry
                bad = [p for p in self._unrecognized_params(e.text)
                       if p in params and p != "action"]
                if page == 1 and bad:
                    for p in bad:
                        params.pop(p)
                    log.info(f"  ⚠  Platform rejected parameter(s) {bad} — "
                             f"dropped, retrying…")
                    page = 0
                    continue
                raise

            if page == 1:
                _debug_save("detection_page1", r.content)

            page_hosts = 0
            for host in root.iter("HOST"):
                hid = host.findtext("ID", default="").strip()
                if not hid:
                    continue
                page_hosts += 1
                hosts_seen.add(hid)
                for det in host.findall("./DETECTION_LIST/DETECTION"):
                    qid = det.findtext("QID", default="").strip()
                    if not qid:
                        continue
                    dtype = det.findtext("TYPE", default="").strip().lower()
                    if dtype == "info":
                        continue
                    if dtype == "potential" and not INCLUDE_POTENTIAL:
                        skipped_potential += 1
                        continue
                    qid_hosts[qid].add(hid)

            log.info(
                f"  Page {page:<4} {page_hosts:>5} host(s)   "
                f"running total: {len(hosts_seen):,} host(s), "
                f"{len(qid_hosts):,} QID(s)"
            )

            # Qualys signals "more data" with a WARNING block holding the next URL
            next_url = root.findtext(".//WARNING/URL")
            url      = next_url.strip() if next_url else None
            time.sleep(0.3)   # polite gap

        return dict(qid_hosts), len(hosts_seen), skipped_potential

    def diagnose_empty(self):
        """Unfiltered test call — tells the user WHY nothing came back."""
        log.info("\n  Running a diagnostic call with NO severity/status filters…")
        try:
            r     = self._request("GET", self.det_url, params={
                "action": "list", "truncation_limit": 5, "show_results": 0,
            })
            root  = self._parse(r)
            hosts = root.findall(".//HOST")
            dets  = sum(len(h.findall("./DETECTION_LIST/DETECTION")) for h in hosts)
        except Exception as exc:                       # diagnostic must never crash
            log.info(f"  ✘  Diagnostic call failed: {exc}")
            return
        if not hosts:
            log.info(
                "  ✘  Even without filters the API returned NO hosts.\n"
                "     → Your account sees no scanned hosts. Check the user's role /\n"
                "       asset-group scope, and that scanners or agents have completed\n"
                "       VM scans (GAV inventory alone does not create detections)."
            )
        else:
            log.info(
                f"  ✔  Unfiltered call returned {len(hosts)} host(s) / {dets} detection(s).\n"
                "     → The API works; your filters are too strict. Review\n"
                "       SEVERITY_LEVELS, DETECTION_STATUS and EXTRA_DETECTION_PARAMS."
            )

    # ── Step 2: Knowledge Base enrichment ────────────────────────────────────
    def fetch_kb(self, qids: list) -> dict:
        """
        Returns { "qid": {title, severity, cves[], cvss2, cvss3, intel[],
                          published, patchable, software} }
        """
        url   = f"{self.base_url}/api/2.0/fo/knowledge_base/vuln/"
        qids  = sorted(qids, key=lambda q: int(q) if q.isdigit() else 0)
        total = len(qids)
        kb    = {}

        for start in range(0, total, KB_BATCH_SIZE):
            batch = qids[start: start + KB_BATCH_SIZE]
            end   = min(start + KB_BATCH_SIZE, total)
            log.info(f"  Knowledge Base [{start+1}–{end} of {total}]…")

            r    = self._request(
                "POST", url,
                data={"action": "list", "ids": ",".join(batch), "details": "All"},
            )
            if start == 0:
                _debug_save("knowledge_base_batch1", r.content)
            root = self._parse(r)

            for v in root.findall(".//VULN_LIST/VULN"):
                qid = v.findtext("QID", default="").strip()
                if not qid:
                    continue
                cves = [
                    c.findtext("ID", default="").strip().upper()
                    for c in v.findall("./CVE_LIST/CVE")
                ]
                # Qualys's own vendor/product tags (informational — shown in QID Mapping)
                software = sorted({
                    f"{s.findtext('VENDOR', default='').strip()} "
                    f"{s.findtext('PRODUCT', default='').strip()}".strip()
                    for s in v.findall("./SOFTWARE_LIST/SOFTWARE")
                } - {""})
                kb[qid] = {
                    "title":     v.findtext("TITLE", default="").strip(),
                    "severity":  int(v.findtext("SEVERITY_LEVEL", default="0") or 0),
                    "cves":      [c for c in cves if c.startswith("CVE-")],
                    "cvss2":     _to_float(v.findtext("./CVSS/BASE")),
                    "cvss3":     _to_float(v.findtext("./CVSS_V3/BASE")),
                    "intel":     [
                        t.text.strip()
                        for t in v.findall("./THREAT_INTELLIGENCE/THREAT_INTEL")
                        if t.text
                    ],
                    "published": (v.findtext("PUBLISHED_DATETIME", default="") or "")[:10],
                    "patchable": v.findtext("PATCHABLE", default="0").strip() == "1",
                    "software":  software,
                }
            time.sleep(0.3)

        return kb

    # ── Step 3a: aggregate by CVE ────────────────────────────────────────────
    @staticmethod
    def build_cve_table(qid_hosts: dict, kb: dict) -> tuple[dict, int]:
        """
        Collapse QID-level data into CVE-level data.
        Asset count = UNION of hosts across every QID that references the CVE.
        """
        cves        = {}
        no_cve_qids = 0

        for qid, hosts in qid_hosts.items():
            meta = kb.get(qid)
            if not meta or not meta["cves"]:
                no_cve_qids += 1      # misconfigs, info gathering, etc.
                continue
            app = group_key(meta["title"])

            for cve in meta["cves"]:
                e = cves.setdefault(cve, {
                    "hosts": set(), "qids": set(), "apps": set(),
                    "title": "", "title_hosts": -1,
                    "severity": 0, "cvss2": None, "cvss3": None,
                    "intel": set(), "published": "", "patchable": False,
                })
                e["hosts"] |= hosts
                e["qids"].add(qid)
                e["apps"].add(app)

                # Use the title of the QID that touches the most hosts
                if len(hosts) > e["title_hosts"]:
                    e["title"], e["title_hosts"] = meta["title"], len(hosts)

                e["severity"]  = max(e["severity"], meta["severity"])
                e["cvss2"]     = _max_none(e["cvss2"], meta["cvss2"])
                e["cvss3"]     = _max_none(e["cvss3"], meta["cvss3"])
                e["intel"].update(meta["intel"])
                e["patchable"] = e["patchable"] or meta["patchable"]
                if meta["published"] and (not e["published"] or meta["published"] < e["published"]):
                    e["published"] = meta["published"]

        return cves, no_cve_qids

    @staticmethod
    def _flags(intel) -> tuple[str, str]:
        low = [t.lower() for t in intel]
        return ("Yes" if any("cisa" in t for t in low) else "No",
                "Yes" if any("active_attacks" in t for t in low) else "No")

    # ── View 1: all CVEs ─────────────────────────────────────────────────────
    def all_cves_frame(self, cves: dict) -> pd.DataFrame:
        rows = []
        for cve, e in cves.items():
            count = len(e["hosts"])
            if count <= MIN_ASSET_COUNT:
                continue
            kev, atk = self._flags(e["intel"])
            apps = sorted(e["apps"])
            rows.append({
                "CVE ID":             cve,
                "Affected Assets":    count,
                "Application(s)":     _clip("; ".join(apps)),
                "QID Count":          len(e["qids"]),
                "QIDs":               _clip(", ".join(sorted(e["qids"], key=int))),
                "Title":              e["title"],
                "Qualys Severity":    e["severity"],
                "CVSS v3":            e["cvss3"],
                "CVSS v2":            e["cvss2"],
                "CISA KEV":           kev,
                "Active Attacks":     atk,
                "Threat Intel Tags":  ", ".join(sorted(e["intel"])),
                "Patchable":          "Yes" if e["patchable"] else "No",
                "Published (Qualys)": e["published"],
            })
        rows.sort(key=lambda x: (-x["Affected Assets"], -(x["CVSS v3"] or 0),
                                 -x["Qualys Severity"], x["CVE ID"]))
        df = pd.DataFrame(rows)
        if not df.empty:
            df.insert(0, "Rank", range(1, len(df) + 1))
        return df

    # ── View 2: applications, each with its top CVEs ─────────────────────────
    @staticmethod
    def build_apps(qid_hosts: dict, kb: dict) -> tuple[dict, list]:
        """
        Club QIDs into applications.
          apps[name] = {"qids": set, "cves": {cve: {"hosts": set, "qids": set}},
                        "titles": {title: hosts}}
        QIDs without a CVE are ignored.  Returns (apps, qid_mapping_rows).
        """
        apps, qid_rows = {}, []
        for qid, hosts in qid_hosts.items():
            meta = kb.get(qid)
            if not meta or not meta["cves"]:
                continue
            name = group_key(meta["title"])
            a    = apps.setdefault(name, {"qids": set(), "cves": {}, "titles": {}})
            a["qids"].add(qid)
            a["titles"][meta["title"]] = max(a["titles"].get(meta["title"], 0), len(hosts))
            for cve in meta["cves"]:
                c = a["cves"].setdefault(cve, {"hosts": set(), "qids": set()})
                c["hosts"] |= hosts
                c["qids"].add(qid)
            qid_rows.append({
                "QID": qid, "Title": meta["title"], "Application (grouped as)": name,
                "Hosts": len(hosts), "CVE Count": len(meta["cves"]),
                "Qualys Vendor / Product": "; ".join(meta["software"][:5]),
            })
        qid_rows.sort(key=lambda r: (r["Application (grouped as)"].lower(), -r["Hosts"]))
        return apps, qid_rows

    def app_frames(self, apps: dict, cves: dict, qid_hosts: dict, kb: dict):
        """
        Returns (by_application_df, top_cves_per_app_df).
        CVEs below MIN_ASSET_COUNT are ignored; applications left with no CVEs are dropped.
        """
        built = []
        for name, a in apps.items():
            recs, all_hosts = [], set()
            for cve, c in a["cves"].items():
                if len(cves[cve]["hosts"]) <= MIN_ASSET_COUNT:
                    continue
                all_hosts |= c["hosts"]
                e        = cves[cve]
                kev, atk = self._flags(e["intel"])
                top_qid  = max(c["qids"], key=lambda q: len(qid_hosts[q]))
                recs.append({
                    "cve": cve, "assets": len(c["hosts"]),
                    "cvss3": e["cvss3"], "cvss2": e["cvss2"],
                    "kev": kev, "atk": atk,
                    "intel": ", ".join(sorted(e["intel"])),
                    "published": e["published"],
                    "title": kb[top_qid]["title"],
                    "qids": ", ".join(sorted(c["qids"], key=int)),
                })
            if not recs:
                continue
            # KEV first → actively attacked → CVSS v3 → reach → id
            recs.sort(key=lambda r: (r["kev"] != "Yes", r["atk"] != "Yes",
                                     -(r["cvss3"] or 0), -r["assets"], r["cve"]))
            built.append({
                "name": name, "a": a, "assets": len(all_hosts), "recs": recs,
                "kev": sum(r["kev"] == "Yes" for r in recs),
                "atk": sum(r["atk"] == "Yes" for r in recs),
                "max3": max((r["cvss3"] for r in recs if r["cvss3"] is not None), default=None),
            })

        built.sort(key=lambda b: (-b["assets"], -b["kev"], -(b["max3"] or 0), b["name"].lower()))

        app_rows, top_rows = [], []
        for rank, b in enumerate(built, start=1):
            recs   = b["recs"]
            shown  = recs[:CVES_PER_APP] if CVES_PER_APP > 0 else recs
            titles = sorted(b["a"]["titles"], key=lambda t: -b["a"]["titles"][t])
            sample = " | ".join(titles[:2]) + (f"  (+{len(titles) - 2} more)" if len(titles) > 2 else "")
            qids   = sorted(b["a"]["qids"], key=int)
            qid_s  = ", ".join(qids[:40]) + (f" … (+{len(qids) - 40} more)" if len(qids) > 40 else "")

            app_rows.append({
                "Rank":                rank,
                "Application":         b["name"],
                "Affected Assets":     b["assets"],
                "QID Count":           len(qids),
                "Total CVEs":          len(recs),
                "Max CVSS v3":         b["max3"],
                "KEV CVEs":            b["kev"],
                "Active-Attack CVEs":  b["atk"],
                f"Top {len(shown)} CVEs" if CVES_PER_APP > 0 else "All CVEs":
                                       _clip(", ".join(r["cve"] for r in shown)),
                "Sample Titles":       sample,
                "QIDs":                qid_s,
            })
            for i, r in enumerate(shown, start=1):
                top_rows.append({
                    "App Rank":            rank,
                    "Application":         b["name"],
                    "App Affected Assets": b["assets"],
                    "CVE # in App":        i,
                    "CVE ID":              r["cve"],
                    "CVE Assets (in app)": r["assets"],
                    "CVSS v3":             r["cvss3"],
                    "CVSS v2":             r["cvss2"],
                    "CISA KEV":            r["kev"],
                    "Active Attacks":      r["atk"],
                    "Threat Intel Tags":   r["intel"],
                    "Published (Qualys)":  r["published"],
                    "Title":               r["title"],
                    "QIDs":                _clip(r["qids"]),
                })
        return pd.DataFrame(app_rows), pd.DataFrame(top_rows)

    # ── Step 4: Excel export ─────────────────────────────────────────────────
    @staticmethod
    def export(sheets: dict, info: list) -> str:
        ts          = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
        base, ext   = os.path.splitext(OUTPUT_FILE)
        output_file = f"{base}_{ts}{ext or '.xlsx'}"

        with pd.ExcelWriter(output_file, engine="openpyxl") as xw:
            for name, frame in sheets.items():
                frame.to_excel(xw, sheet_name=name, index=False)
            pd.DataFrame(info, columns=["Parameter", "Value"]).to_excel(
                xw, sheet_name="Run Info", index=False)

        _apply_styles(output_file)
        return output_file

    # ── Console preview ──────────────────────────────────────────────────────
    @staticmethod
    def _preview(cve_df: pd.DataFrame, app_df: pd.DataFrame):
        log.info("\n  TOP 10 APPLICATIONS (by affected assets)")
        log.info(f"  {'#':<4} {'Application':<36} {'Assets':>7} {'CVEs':>5} {'KEV':>4}")
        log.info(f"  {'─'*4} {'─'*36} {'─'*7} {'─'*5} {'─'*4}")
        for _, r in app_df.head(10).iterrows():
            log.info(f"  {r['Rank']:<4} {str(r['Application'])[:36]:<36} "
                     f"{r['Affected Assets']:>7} {r['Total CVEs']:>5} {r['KEV CVEs']:>4}")

        log.info("\n  TOP 10 CVEs (by affected assets)")
        log.info(f"  {'#':<4} {'CVE ID':<16} {'Assets':>7} {'CVSS3':>5} {'KEV':>4}  Application")
        log.info(f"  {'─'*4} {'─'*16} {'─'*7} {'─'*5} {'─'*4}  {'─'*24}")
        for _, r in cve_df.head(10).iterrows():
            c3 = f"{r['CVSS v3']:.1f}" if pd.notna(r["CVSS v3"]) else "-"
            log.info(f"  {r['Rank']:<4} {r['CVE ID']:<16} {r['Affected Assets']:>7} "
                     f"{c3:>5} {r['CISA KEV']:>4}  {str(r['Application(s)'])[:24]}")

    # ── Main flow ─────────────────────────────────────────────────────────────
    def run(self):
        self._check_config()

        log.info(DLINE)
        log.info("  Qualys Critical Vulnerability Watchlist Builder  (VMDR)")
        log.info(DLINE)
        log.info(f"  Severity levels     : {SEVERITY_LEVELS}")
        log.info(f"  Detection type      : {'confirmed + potential' if INCLUDE_POTENTIAL else 'confirmed only'}")
        log.info(f"  Detection status    : {DETECTION_STATUS}")
        log.info(f"  CVEs per application: {CVES_PER_APP if CVES_PER_APP > 0 else 'all'}")
        log.info(f"  Min asset count     : {'> ' + str(MIN_ASSET_COUNT) if MIN_ASSET_COUNT > 0 else 'none (all CVEs)'}")
        log.info(LINE)

        self.login()

        try:
            # ── STEP 1 — FETCH DETECTIONS ─────────────────────────────────────
            log.info("\n  STEP 1 — Fetching detections (this can take a while)…\n")
            qid_hosts, hosts_returned, skipped_potential = self.fetch_detections()

            if not qid_hosts:
                log.info("\n  ✘  No usable detections found.")
                if hosts_returned == 0:
                    self.diagnose_empty()
                else:
                    log.info(
                        f"     {hosts_returned:,} host(s) were returned, but every detection was\n"
                        f"     a Potential one ({skipped_potential:,} skipped). "
                        f"Set INCLUDE_POTENTIAL = True to count them."
                    )
                return

            host_count = len(set().union(*qid_hosts.values()))
            log.info(
                f"\n  ✔  {host_count:,} host(s) with {len(qid_hosts):,} "
                f"unique QID(s) matched your filters"
                + (f"  ({skipped_potential:,} potential detection(s) skipped)"
                   if skipped_potential else "")
            )

            # ── STEP 2 — ENRICH ───────────────────────────────────────────────
            log.info("\n  STEP 2 — Enriching QIDs from the Knowledge Base…\n")
            kb = self.fetch_kb(list(qid_hosts.keys()))
            log.info(f"\n  ✔  {len(kb):,} / {len(qid_hosts):,} QID(s) resolved")
            if not kb:
                raise RuntimeError(
                    "The Knowledge Base returned no data for any QID. "
                    "Check the account has access to the Knowledge Base API "
                    "(re-run with DEBUG_SAVE_XML = True to inspect the response)."
                )

            # ── STEP 3 — BUILD THE TWO VIEWS ──────────────────────────────────
            cves, no_cve_qids = self.build_cve_table(qid_hosts, kb)
            if not cves:
                log.info("\n  ✘  None of the detected QIDs map to a CVE. Nothing to export.")
                return

            cve_df               = self.all_cves_frame(cves)
            apps, qid_rows       = self.build_apps(qid_hosts, kb)
            app_df, top_cves_df  = self.app_frames(apps, cves, qid_hosts, kb)

            log.info("")
            log.info(LINE)
            log.info("  RESULT")
            log.info(LINE)
            log.info(f"  Unique CVEs found                : {len(cves):,}")
            log.info(f"  CVEs exported                    : {len(cve_df):,}")
            log.info(f"  Applications (grouped)           : {len(app_df):,}")
            log.info(f"  QIDs without a CVE (skipped)     : {no_cve_qids:,}")

            if cve_df.empty:
                top = sorted(cves.items(), key=lambda kv: -len(kv[1]["hosts"]))[:10]
                log.info(f"\n  ✘  No CVE affects more than {MIN_ASSET_COUNT} host(s). "
                         f"Highest counts seen:\n")
                for cve, e in top:
                    log.info(f"  {cve:<16} {len(e['hosts']):>7}  {e['title'][:45]}")
                log.info("\n  Lower MIN_ASSET_COUNT (or set it to 0) and re-run. No file written.")
                return

            self._preview(cve_df, app_df)

            # ── STEP 4 — EXPORT ───────────────────────────────────────────────
            kev_total = int((cve_df["CISA KEV"] == "Yes").sum())
            info = [
                ("Generated (UTC)",                  datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S")),
                ("Qualys platform",                  self.base_url),
                ("Severity levels",                  SEVERITY_LEVELS),
                ("Detection type",                   "confirmed + potential" if INCLUDE_POTENTIAL else "confirmed only"),
                ("Detection status",                 DETECTION_STATUS),
                ("Extra detection filters",          str(EXTRA_DETECTION_PARAMS) if EXTRA_DETECTION_PARAMS else "none"),
                ("CVES_PER_APP (0 = all)",           CVES_PER_APP),
                ("MIN_ASSET_COUNT (0 = none)",       MIN_ASSET_COUNT),
                ("PRODUCT_OVERRIDES rules",          len(PRODUCT_OVERRIDES)),
                ("Hosts with matching detections",   host_count),
                ("Unique QIDs matched",              len(qid_hosts)),
                ("Potential detections skipped",     skipped_potential),
                ("QIDs without a CVE (skipped)",     no_cve_qids),
                ("CVEs exported",                    len(cve_df)),
                ("CVEs on CISA KEV list",            kev_total),
                ("Applications (grouped)",           len(app_df)),
            ]
            sheets = {
                "All CVEs":         cve_df,
                "By Application":   app_df,
                "Top CVEs per App": top_cves_df,
                "QID Mapping":      pd.DataFrame(qid_rows),
            }
            output_file = self.export(sheets, info)

            log.info("")
            log.info(DLINE)
            log.info("  COMPLETE")
            log.info(DLINE)
            log.info(f"  Results saved → {output_file}\n")

        finally:
            self.logout()


# ── Excel styling ─────────────────────────────────────────────────────────────
def _apply_styles(filepath: str):
    wb     = load_workbook(filepath)
    thin   = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    RED    = PatternFill("solid", fgColor="FFC7CE")
    AMBER  = PatternFill("solid", fgColor="FFEB9C")
    ALT    = PatternFill("solid", fgColor="EBF3FB")
    head_font, body_font = (Font(bold=True, color="FFFFFF", name="Calibri", size=11),
                            Font(name="Calibri", size=10))
    head_fill = PatternFill("solid", fgColor="1F4E79")
    head_align, body_align = (Alignment(horizontal="center", vertical="center"),
                              Alignment(horizontal="left", vertical="center"))

    is_yes = lambda v: v == "Yes"
    is_pos = lambda v: isinstance(v, (int, float)) and v > 0
    # header name → (test, fill)
    marks = {
        "CISA KEV":           (is_yes, RED),
        "Active Attacks":     (is_yes, AMBER),
        "KEV CVEs":           (is_pos, RED),
        "Active-Attack CVEs": (is_pos, AMBER),
    }

    for ws in wb.worksheets:
        for cell in ws[1]:
            cell.fill, cell.font, cell.alignment, cell.border = head_fill, head_font, head_align, border

        mark_cols = {c.column: marks[c.value] for c in ws[1] if c.value in marks}

        for row_idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
            for cell in row:
                cell.font, cell.border, cell.alignment = body_font, border, body_align
                if row_idx % 2 == 0:
                    cell.fill = ALT
                if cell.column in mark_cols:
                    test, fill = mark_cols[cell.column]
                    if test(cell.value):
                        cell.fill = fill

        # Auto column widths
        for col in ws.columns:
            letter    = get_column_letter(col[0].column)
            max_width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
            ws.column_dimensions[letter].width = min(max_width + 4, 55)

        ws.freeze_panes    = "A2"
        ws.auto_filter.ref = ws.dimensions

    wb.save(filepath)


# ── Entry point ───────────────────────────────────────────────────────────────
if __name__ == "__main__":
    try:
        QualysCVEWatchlist().run()
    except requests.exceptions.SSLError as exc:
        log.info(f"\n  ✘  SSL error: {exc}\n     Check CERT_PATH (or your corporate proxy certificate).\n")
        sys.exit(1)
    except requests.exceptions.RequestException as exc:
        log.info(f"\n  ✘  Network error: {exc}\n     Check BASE_URL and connectivity.\n")
        sys.exit(1)
    except (ValueError, RuntimeError) as exc:      # includes QualysAPIError
        log.info(f"\n  ✘  {exc}\n")
        sys.exit(1)
